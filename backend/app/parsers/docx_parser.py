from pathlib import Path
from typing import List
import docx
from app.parsers.base import BaseParser, ParsedDocument

class DocxParser(BaseParser):
    async def parse(self, file_path: Path, filename: str) -> ParsedDocument:
        doc = docx.Document(str(file_path))
        text_parts: List[str] = []
        headings: List[str] = []
        
        # Parse paragraphs
        for para in doc.paragraphs:
            text = para.text.strip()
            if not text:
                continue
            if para.style and para.style.name.startswith("Heading"):
                headings.append(text)
                text_parts.append(f"\n## {text}\n")
            else:
                text_parts.append(text)
        
        # Parse tables into markdown format
        for t_idx, table in enumerate(doc.tables):
            text_parts.append(f"\n[Table {t_idx + 1}]")
            table_rows = []
            for row in table.rows:
                row_cells = [cell.text.strip().replace("\n", " ") for cell in row.cells]
                table_rows.append(row_cells)
            
            if table_rows:
                header = table_rows[0]
                text_parts.append("| " + " | ".join(header) + " |")
                text_parts.append("| " + " | ".join(["---"] * len(header)) + " |")
                for data_row in table_rows[1:]:
                    text_parts.append("| " + " | ".join(data_row) + " |")
            text_parts.append("")
        
        full_text = "\n".join(text_parts)
        token_estimate = max(1, len(full_text) // 4)
        
        summary = f"DOCX document with {len(doc.paragraphs)} paragraphs and {len(doc.tables)} tables."
        if headings:
            summary += f" Key sections: {', '.join(headings[:8])}"
        
        return ParsedDocument(
            filename=filename,
            file_type="docx",
            text_content=full_text,
            token_estimate=token_estimate,
            summary_or_toc=summary,
            metadata={"headings": headings, "table_count": len(doc.tables)}
        )
